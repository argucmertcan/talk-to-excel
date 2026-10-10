"""Demo veri üreticisi: Nova Elektronik 2025 satış dosyası.

Rastgele değil, içine bilinçli hikâyeler gömülü:
  1. Ankara Q3 düşüşü        — Ankara Kızılay mağazası Ağu–Eyl tadilatta, satış yok
  2. Akıllı Saat X2 iadeleri  — iade oranı ~%27 (ortalama ~%4), memnuniyet düşük
  3. Temsilci uçları          — Elif Yıldız hedefin ~%140'ı, Burak Şahin ~%60'ı
  4. Kasım kampanyası         — ciro ~2 kat, indirim yüksek, marj düşük
  5. Online yükselişi         — online payı Ocak ~%22 → Aralık ~%42
  6. Zarar eden siparişler    — Beyaz Eşya'da %35+ indirimli siparişler zararda
  7. Veri kalitesi            — 3 mükerrer sipariş, "Istanbul" yazım hatası, boş memnuniyet

Çalıştır:  uv run --with openpyxl python demo/generate_demo.py
"""
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

random.seed(2025)
OUT = Path(__file__).parent / "nova_elektronik_satis_2025.xlsx"

REGIONS = {
    "Marmara": {"İstanbul": ["Kadıköy", "Beşiktaş", "Bakırköy"], "Bursa": ["Nilüfer"], "Kocaeli": ["İzmit"]},
    "Ege": {"İzmir": ["Alsancak", "Karşıyaka"], "Manisa": ["Merkez"]},
    "İç Anadolu": {"Ankara": ["Kızılay", "Çankaya"], "Konya": ["Selçuklu"], "Eskişehir": ["Tepebaşı"]},
    "Akdeniz": {"Antalya": ["Muratpaşa"], "Adana": ["Seyhan"], "Mersin": ["Yenişehir"]},
    "Karadeniz": {"Samsun": ["Atakum"], "Trabzon": ["Ortahisar"]},
    "Doğu & Güneydoğu": {"Gaziantep": ["Şehitkamil"], "Diyarbakır": ["Kayapınar"], "Erzurum": ["Yakutiye"]},
}
REGION_WEIGHT = {"Marmara": 34, "Ege": 16, "İç Anadolu": 18, "Akdeniz": 14, "Karadeniz": 8, "Doğu & Güneydoğu": 10}
CITY_WEIGHT = {"İstanbul": 70, "Bursa": 18, "Kocaeli": 12, "İzmir": 75, "Manisa": 25, "Ankara": 62,
               "Konya": 20, "Eskişehir": 18, "Antalya": 45, "Adana": 32, "Mersin": 23, "Samsun": 55,
               "Trabzon": 45, "Gaziantep": 45, "Diyarbakır": 30, "Erzurum": 25}

# ürün kodu, ad, kategori, liste fiyatı, maliyet oranı, tedarikçi
PRODUCTS = [
    ("TV-55", '55" 4K Televizyon', "Televizyon", 28999, 0.74, "Ekran Dağıtım"),
    ("TV-65", '65" QLED Televizyon', "Televizyon", 46999, 0.72, "Ekran Dağıtım"),
    ("TV-43", '43" Full HD Televizyon', "Televizyon", 14999, 0.78, "Ekran Dağıtım"),
    ("PC-14", "14\" Ultrabook", "Bilgisayar", 38999, 0.80, "Bilişim Tedarik"),
    ("PC-16", "16\" Oyun Laptopu", "Bilgisayar", 61999, 0.79, "Bilişim Tedarik"),
    ("PC-TB", "11\" Tablet", "Bilgisayar", 12999, 0.76, "Bilişim Tedarik"),
    ("PC-MN", "27\" Monitör", "Bilgisayar", 8999, 0.70, "Bilişim Tedarik"),
    ("TL-PR", "Akıllı Telefon Pro", "Telefon", 54999, 0.83, "Mobil Grup"),
    ("TL-ST", "Akıllı Telefon Standart", "Telefon", 24999, 0.81, "Mobil Grup"),
    ("TL-EK", "Akıllı Telefon Ekonomik", "Telefon", 9999, 0.79, "Mobil Grup"),
    ("BE-BZ", "No-Frost Buzdolabı", "Beyaz Eşya", 32999, 0.71, "Ev Sanayi"),
    ("BE-CM", "Çamaşır Makinesi 9 kg", "Beyaz Eşya", 21999, 0.70, "Ev Sanayi"),
    ("BE-BL", "Bulaşık Makinesi", "Beyaz Eşya", 18999, 0.70, "Ev Sanayi"),
    ("BE-KR", "Kurutma Makinesi", "Beyaz Eşya", 24999, 0.69, "Ev Sanayi"),
    ("KE-SP", "Robot Süpürge", "Küçük Ev Aletleri", 13999, 0.62, "Ev Sanayi"),
    ("KE-KM", "Espresso Makinesi", "Küçük Ev Aletleri", 9499, 0.60, "Mutfak Dünyası"),
    ("KE-AF", "Airfryer XL", "Küçük Ev Aletleri", 4999, 0.58, "Mutfak Dünyası"),
    ("KE-BL", "Blender Seti", "Küçük Ev Aletleri", 2799, 0.55, "Mutfak Dünyası"),
    ("SS-KK", "Kablosuz Kulaklık", "Ses Sistemleri", 5499, 0.57, "Ses Teknik"),
    ("SS-SB", "Soundbar 5.1", "Ses Sistemleri", 15999, 0.63, "Ses Teknik"),
    ("SS-HP", "Bluetooth Hoparlör", "Ses Sistemleri", 2999, 0.54, "Ses Teknik"),
    ("GY-X2", "Akıllı Saat X2", "Giyilebilir", 7999, 0.61, "Mobil Grup"),
    ("GY-BL", "Akıllı Bileklik", "Giyilebilir", 1999, 0.52, "Mobil Grup"),
    ("GY-KL", "Spor Kulaklık", "Giyilebilir", 3499, 0.55, "Ses Teknik"),
]
CAT_WEIGHT = {"Televizyon": 14, "Bilgisayar": 16, "Telefon": 20, "Beyaz Eşya": 15,
              "Küçük Ev Aletleri": 16, "Ses Sistemleri": 10, "Giyilebilir": 9}

REPS = {  # temsilci: (bölge, yıllık hedef ₺, performans çarpanı, başlangıç)
    "Elif Yıldız": ("Marmara", 0, 1.55, date(2021, 3, 1)),
    "Mert Kaya": ("Marmara", 0, 1.0, date(2019, 9, 16)),
    "Zeynep Arslan": ("Marmara", 0, 0.95, date(2023, 1, 9)),
    "Can Öztürk": ("Ege", 0, 1.05, date(2020, 5, 4)),
    "Deniz Aydın": ("Ege", 0, 0.92, date(2022, 7, 18)),
    "Burak Şahin": ("İç Anadolu", 0, 0.62, date(2024, 2, 5)),
    "Selin Koç": ("İç Anadolu", 0, 1.08, date(2018, 11, 12)),
    "Emre Çelik": ("Akdeniz", 0, 1.0, date(2021, 6, 21)),
    "Ayşe Demir": ("Akdeniz", 0, 0.97, date(2020, 1, 13)),
    "Oğuz Polat": ("Karadeniz", 0, 1.02, date(2019, 4, 1)),
    "Gizem Kurt": ("Karadeniz", 0, 0.94, date(2023, 9, 4)),
    "Hakan Yılmaz": ("Doğu & Güneydoğu", 0, 1.03, date(2017, 10, 2)),
    "Nazlı Erdem": ("Doğu & Güneydoğu", 0, 0.96, date(2022, 3, 28)),
}
MONTH_FACTOR = {1: 0.82, 2: 0.78, 3: 0.92, 4: 0.95, 5: 1.0, 6: 1.05, 7: 0.98,
                8: 0.96, 9: 1.0, 10: 1.05, 11: 2.05, 12: 1.45}
MONTHS_TR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz",
             "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
BASE_PER_MONTH = 420


def wpick(d):
    keys = list(d)
    return random.choices(keys, weights=[d[k] for k in keys])[0]


def online_share(month):  # Ocak %22 → Aralık %42
    return 0.22 + (month - 1) * (0.20 / 11)


rows = []
order_no = 250000
for month in range(1, 13):
    n = int(BASE_PER_MONTH * MONTH_FACTOR[month] * random.uniform(0.96, 1.04))
    days_in_month = (date(2025 + month // 12, month % 12 + 1, 1) - date(2025, month, 1)).days
    for _ in range(n):
        region = wpick(REGION_WEIGHT)
        cities = REGIONS[region]
        city = random.choices(list(cities), weights=[CITY_WEIGHT[c] for c in cities])[0]
        store = f"{city} {random.choice(cities[city])}"
        # Hikâye 1: Ankara Kızılay Ağu–Eyl tadilatta; Temmuz'da da zayıf
        if store == "Ankara Kızılay" and month in (8, 9):
            continue
        if city == "Ankara" and month == 7 and random.random() < 0.35:
            continue
        rep_pool = [r for r, v in REPS.items() if v[0] == region]
        rep = random.choices(rep_pool, weights=[REPS[r][2] for r in rep_pool])[0]

        channel_r = random.random()
        if channel_r < online_share(month):
            channel = "Online"
        elif channel_r < online_share(month) + 0.12:
            channel = "Bayi"
        else:
            channel = "Mağaza"
        if channel == "Online":
            store = "Online Mağaza"

        cat = wpick(CAT_WEIGHT)
        prod = random.choice([p for p in PRODUCTS if p[2] == cat])
        code, name, _, list_price, cost_ratio, _ = prod

        corporate = random.random() < 0.18
        qty = random.choices([1, 2, 3, 4, 5], weights=[70, 16, 7, 4, 3])[0]
        if corporate:
            qty = random.choice([2, 3, 5, 8, 10, 12, 15])

        # İndirim: normal ay %0–15, Kasım %10–30; Beyaz Eşya Kasım kampanyasında %35–45 uçlar
        if month == 11:
            disc = random.choice([10, 15, 20, 20, 25, 30])
            if cat == "Beyaz Eşya" and random.random() < 0.45:
                disc = random.choice([35, 38, 40, 42, 45])
        elif month == 12:
            disc = random.choice([0, 5, 10, 10, 15, 20])
        else:
            disc = random.choices([0, 5, 10, 15], weights=[45, 25, 20, 10])[0]
        if corporate:
            disc = min(disc + 5, 45)

        unit = round(list_price * random.uniform(0.97, 1.03), -1)
        revenue = round(unit * qty * (1 - disc / 100), 2)
        cost = round(list_price * cost_ratio * qty * random.uniform(0.98, 1.02), 2)
        profit = round(revenue - cost, 2)

        # Hikâye 2: Akıllı Saat X2 iadesi ~%27
        ret_p = 0.27 if code == "GY-X2" else 0.035 + (0.02 if channel == "Online" else 0)
        returned = random.random() < ret_p
        if code == "GY-X2":
            sat = random.choices([1, 2, 3, 4, 5], weights=[30, 30, 20, 12, 8])[0]
        else:
            sat = random.choices([1, 2, 3, 4, 5], weights=[3, 5, 14, 38, 40])[0]
        if returned:
            sat = min(sat, random.choice([1, 2, 3]))

        d = date(2025, month, random.randint(1, days_in_month))
        order_no += random.randint(1, 3)
        rows.append({
            "Sipariş No": f"NV-{order_no}", "Tarih": d, "Bölge": region, "Şehir": city,
            "Mağaza": store, "Kanal": channel, "Kategori": cat, "Ürün Kodu": code, "Ürün": name,
            "Adet": qty, "Birim Fiyat (₺)": unit, "İndirim (%)": disc, "Ciro (₺)": revenue,
            "Maliyet (₺)": cost, "Kâr (₺)": profit, "Satış Temsilcisi": rep,
            "Müşteri Tipi": "Kurumsal" if corporate else "Bireysel",
            "Ödeme": random.choices(["Kredi Kartı", "Taksit", "Havale/EFT"], weights=[55, 32, 13])[0],
            "İade": "Evet" if returned else "Hayır", "Memnuniyet (1-5)": sat,
        })

rows.sort(key=lambda r: (r["Tarih"], r["Sipariş No"]))

# Hikâye 3: hedefler gerçekleşenden türetilir, böylece Elif ~%140, Burak ~%60 olur
actual = {r: 0.0 for r in REPS}
for r in rows:
    actual[r["Satış Temsilcisi"]] += r["Ciro (₺)"]
target_ratio = {"Elif Yıldız": 1.41, "Burak Şahin": 0.61}
targets = {}
for rep, total in actual.items():
    ratio = target_ratio.get(rep, random.uniform(0.92, 1.12))
    targets[rep] = round(total / ratio, -5)

# Hikâye 7: veri kalitesi sorunları
for r in random.sample([r for r in rows if r["Şehir"] == "İstanbul"], 14):
    r["Şehir"] = "Istanbul"
for r in random.sample(rows, 23):
    r["Memnuniyet (1-5)"] = None
for r in random.sample(rows, 3):
    rows.insert(rows.index(r) + 1, dict(r))

# ---------- Yazım ----------
wb = Workbook()
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
THIN = Side(style="thin", color="D0D7E2")
TL = "#,##0.00 ₺"


def write_sheet(ws, headers, data, widths, formats=None):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    for row in data:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if formats:
        for col, fmt in formats.items():
            for cell in ws[col][1:]:
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


ws = wb.active
ws.title = "Satışlar"
headers = list(rows[0])
write_sheet(ws, headers, [[r[h] for h in headers] for r in rows],
            [12, 11, 17, 11, 20, 9, 18, 10, 24, 6, 15, 10, 15, 15, 14, 16, 12, 12, 7, 11],
            {"B": "dd.mm.yyyy", "K": TL, "M": TL, "N": TL, "O": TL})

ws = wb.create_sheet("Ürünler")
stock = {p[0]: random.randint(15, 420) for p in PRODUCTS}
stock["GY-X2"] = 612  # iade dönüşleriyle şişmiş stok
write_sheet(ws, ["Ürün Kodu", "Ürün", "Kategori", "Liste Fiyatı (₺)", "Birim Maliyet (₺)", "Tedarikçi", "Stok"],
            [[p[0], p[1], p[2], p[3], round(p[3] * p[4], 2), p[5], stock[p[0]]] for p in PRODUCTS],
            [11, 26, 18, 16, 17, 16, 8], {"D": TL, "E": TL})

ws = wb.create_sheet("Temsilciler")
write_sheet(ws, ["Satış Temsilcisi", "Bölge", "Yıllık Hedef (₺)", "İşe Başlama"],
            [[rep, v[0], targets[rep], v[3]] for rep, v in REPS.items()],
            [18, 18, 18, 13], {"C": TL, "D": "dd.mm.yyyy"})

ws = wb.create_sheet("Mağazalar")
managers = ["Serkan Aksoy", "Pınar Güneş", "Tolga Yavuz", "Ebru Kılıç", "Murat Aslan", "Sevgi Uçar",
            "Kerem Özdemir", "Derya Şen", "Volkan Tekin", "İpek Bulut", "Onur Kaplan", "Melis Avcı",
            "Cem Doğan", "Seda Karaca", "Barış Ateş", "Hande Çınar", "Ali Rıza Er", "Tuba Sezer",
            "Umut Bal", "Esra Keskin", "Furkan Işık"]
store_rows, i = [], 0
for region, cities in REGIONS.items():
    for city, districts in cities.items():
        for dist in districts:
            note = "Ağu–Eyl 2025 tadilat nedeniyle kapalı" if f"{city} {dist}" == "Ankara Kızılay" else ""
            store_rows.append([f"{city} {dist}", city, region,
                               date(random.randint(2012, 2022), random.randint(1, 12), 1),
                               random.choice([450, 600, 750, 900, 1200]), managers[i % len(managers)], note])
            i += 1
store_rows.append(["Online Mağaza", "—", "Tümü", date(2016, 4, 1), None, "E-ticaret Ekibi", ""])
write_sheet(ws, ["Mağaza", "Şehir", "Bölge", "Açılış", "Alan (m²)", "Mağaza Müdürü", "Not"],
            store_rows, [20, 11, 18, 11, 10, 16, 36], {"D": "dd.mm.yyyy"})

ws = wb.create_sheet("Aylık Hedefler")
monthly = {}
for r in rows:
    k = (r["Tarih"].month, r["Bölge"])
    monthly[k] = monthly.get(k, 0) + r["Ciro (₺)"]
target_rows = []
for m in range(1, 13):
    for region in REGIONS:
        base = sum(monthly.get((mm, region), 0) for mm in range(1, 13)) / 12 * MONTH_FACTOR[m] / 1.08
        target_rows.append([MONTHS_TR[m - 1], region, round(base * random.uniform(0.95, 1.05), -4)])
write_sheet(ws, ["Ay", "Bölge", "Hedef Ciro (₺)"], target_rows, [10, 18, 18], {"C": TL})

wb.save(OUT)

# Hızlı doğrulama çıktısı
tot = sum(r["Ciro (₺)"] for r in rows)
print(f"{OUT.name}: {len(rows)} satır, toplam ciro {tot:,.0f} ₺")
for rep in ("Elif Yıldız", "Burak Şahin"):
    print(f"  {rep}: hedefin %{actual[rep] / targets[rep] * 100:.0f}'i")
x2 = [r for r in rows if r["Ürün Kodu"] == "GY-X2"]
print(f"  Akıllı Saat X2 iade: %{sum(r['İade'] == 'Evet' for r in x2) / len(x2) * 100:.0f}")
print(f"  Zarar eden sipariş: {sum(r['Kâr (₺)'] < 0 for r in rows)}")
for q, ms in (("Q2", (4, 5, 6)), ("Q3", (7, 8, 9))):
    print(f"  Ankara {q}: {sum(r['Ciro (₺)'] for r in rows if r['Şehir'] == 'Ankara' and r['Tarih'].month in ms):,.0f} ₺")
